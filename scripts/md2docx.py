#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""markdown → Word（支持 #/##/### 标题、列表、引用、**粗体**、markdown 表格）。"""
import argparse, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

EA_BODY = '宋体'
EA_HEAD = '微软雅黑'
ASCII_BODY = 'Times New Roman'
ASCII_HEAD = 'Microsoft YaHei'
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x59, 0x59, 0x59)


def set_font(run, ea, ascii_f, size, bold=False, color=None):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), ea)
    rFonts.set(qn('w:ascii'), ascii_f)
    rFonts.set(qn('w:hAnsi'), ascii_f)


def add_runs(paragraph, text, ea=EA_BODY, ascii_f=ASCII_BODY, size=10.5,
             base_bold=False, base_color=None):
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_font(run, ea, ascii_f, size, bold=True, color=base_color)
        else:
            run = paragraph.add_run(part)
            set_font(run, ea, ascii_f, size, bold=base_bold, color=base_color)


def cell_text(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def is_sep_row(cells):
    return all(re.fullmatch(r':?-{2,}:?', c) for c in cells)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True, help="输入 markdown 文件")
    ap.add_argument("--out", required=True, help="输出 docx 文件")
    args = ap.parse_args()

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)

    normal = doc.styles['Normal']
    normal.font.name = ASCII_BODY
    normal.font.size = Pt(10.5)
    normal_rpr = normal.element.get_or_add_rPr()
    normal_rfonts = normal_rpr.get_or_add_rFonts()
    normal_rfonts.set(qn('w:eastAsia'), EA_BODY)
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.12

    HEAD_SIZE = {1: 20, 2: 15, 3: 12.5}
    HEAD_BEFORE = {1: 0, 2: 14, 3: 9}
    HEAD_AFTER = {1: 6, 2: 5, 3: 2}

    def add_heading(text, level):
        p = doc.add_heading('', level=level)
        add_runs(p, text, ea=EA_HEAD, ascii_f=ASCII_HEAD,
                 size=HEAD_SIZE[level], base_bold=True, base_color=BLACK)
        pf = p.paragraph_format
        pf.space_before = Pt(HEAD_BEFORE[level])
        pf.space_after = Pt(HEAD_AFTER[level])
        return p

    lines = open(args.md, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        s = line.strip()
        # 跳过 HTML 锚点标签（markdown 里的跳转锚点，Word 里不需要）
        if s.startswith('<a ') and s.endswith('</a>'):
            i += 1
            continue
        if s.startswith('<a id='):
            i += 1
            continue
        # 表格块
        if s.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            header_cells = None
            rows = []
            for b in block:
                cells = cell_text(b)
                if is_sep_row(cells):
                    continue
                if header_cells is None:
                    header_cells = cells
                else:
                    rows.append(cells)
            if header_cells is not None:
                ncols = len(header_cells)
                table = doc.add_table(rows=1, cols=ncols)
                table.style = 'Table Grid'
                table.autofit = True
                hdr = table.rows[0].cells
                for c in range(ncols):
                    hdr[c].text = ''
                    add_runs(hdr[c].paragraphs[0], header_cells[c], size=9, base_bold=True)
                for r in rows:
                    cells = table.add_row().cells
                    for c in range(min(ncols, len(r))):
                        cells[c].text = ''
                        add_runs(cells[c].paragraphs[0], r[c], size=9)
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(2)
            continue
        if s.startswith('### '):
            add_heading(s[4:], 3)
        elif s.startswith('## '):
            add_heading(s[3:], 2)
        elif s.startswith('# '):
            p = add_heading(s[2:], 1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith('> '):
            p = doc.add_paragraph()
            add_runs(p, s[2:], size=9, base_color=GRAY)
            p.paragraph_format.space_after = Pt(4)
        elif s.startswith('---'):
            pass
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, s[2:])
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
        else:
            p = doc.add_paragraph()
            add_runs(p, s)
        i += 1

    doc.save(args.out)
    print(f"[md2docx] 已生成 {args.out}")


if __name__ == "__main__":
    main()
